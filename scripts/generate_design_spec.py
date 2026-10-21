#!/usr/bin/env python3
"""Generate the software design specification from a confirmed business model."""

from __future__ import annotations

import argparse
import platform
import re
import shutil
import subprocess
from copy import deepcopy
from pathlib import Path
from typing import Any

from common import ensure_dir, read_json, safe_filename
from document_structure import DESIGN_CHAPTERS, section_markdown

try:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
except Exception as exc:  # pragma: no cover - environment gate handles this
    Document = None  # type: ignore[assignment]
    OxmlElement = None  # type: ignore[assignment]
    qn = None  # type: ignore[assignment]
    DOCX_IMPORT_ERROR = exc
else:
    DOCX_IMPORT_ERROR = None


REQUIRED_FIELDS = tuple((field, title) for field, title, _pages in DESIGN_CHAPTERS)


def paragraphs(value: Any) -> list[str]:
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    text = str(value or "").strip()
    return [text] if text else []


def section_payload(value: Any, field: str) -> tuple[list[str], list[dict[str, Any]]]:
    if isinstance(value, dict):
        body = paragraphs(value.get("paragraphs") or value.get("body") or value.get("overview"))
        raw_subsections = value.get("subsections") or value.get("sections") or []
    else:
        body = paragraphs(value)
        raw_subsections = []
    subsections: list[dict[str, Any]] = []
    if isinstance(raw_subsections, list):
        for item in raw_subsections:
            if isinstance(item, dict) and str(item.get("title") or "").strip():
                subsections.append(item)
    if not body and not subsections:
        raise SystemExit(
            "STOP_FOR_USER\n"
            f"NEXT_ACTION: 业务理解缺少 `{field}`，无法生成软件设计说明书的对应章节。"
        )
    return body, subsections


def validate_design_structure(business: dict[str, Any]) -> None:
    for field, _title in REQUIRED_FIELDS:
        section_payload(business.get(field), field)
    _body, architecture = section_payload(business.get("design_architecture"), "design_architecture")
    if not architecture:
        raise SystemExit(
            "STOP_FOR_USER\n"
            "NEXT_ACTION: 系统总体架构设计须包含真实分层、组件协作或部署关系小节；接口与数据设计独立成章。"
        )
    function_overview, functions = section_payload(business.get("design_functions"), "design_functions")
    if not function_overview or not functions:
        raise SystemExit(
            "STOP_FOR_USER\n"
            "NEXT_ACTION: 系统功能与模块详细设计须先有总览，再用 subsections 分项说明真实模块。"
        )
    _use_case_overview, use_cases = section_payload(business.get("design_use_cases"), "design_use_cases")
    if not use_cases:
        raise SystemExit(
            "STOP_FOR_USER\n"
            "NEXT_ACTION: 系统关键用例描述须至少包含一个典型功能用例小节。"
        )
    if len(use_cases) < 2 or "典型流程与应用场景" not in str(use_cases[-1].get("title") or ""):
        raise SystemExit("STOP_FOR_USER\nNEXT_ACTION: 第10章须包含关键用例，并以典型流程与应用场景小节结束。")


def table_payload(raw: Any) -> tuple[list[str], list[list[str]]] | None:
    if not isinstance(raw, dict):
        return None
    headers = raw.get("headers")
    rows = raw.get("rows")
    if not isinstance(headers, list) or not headers or not isinstance(rows, list):
        return None
    clean_headers = [str(x) for x in headers]
    clean_rows = [[str(x) for x in row] for row in rows if isinstance(row, list)]
    return clean_headers, clean_rows


def add_table(document: Any, payload: tuple[list[str], list[list[str]]]) -> None:
    headers, rows = payload
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row[: len(headers)]):
            cells[index].text = value


def add_body_paragraph(document: Any, text: str) -> None:
    paragraph = document.add_paragraph(style="正文内容")
    paragraph.add_run(text)


def add_section(document: Any, heading: str, value: Any, field: str, level: int = 1) -> None:
    body, subsections = section_payload(value, field)
    document.add_heading(heading, level=level)
    for paragraph in body:
        add_body_paragraph(document, paragraph)
    if isinstance(value, dict):
        for raw in value.get("tables") or ([value["table"]] if value.get("table") else []):
            payload = table_payload(raw)
            if payload:
                add_table(document, payload)
    for index, subsection in enumerate(subsections, start=1):
        title = str(subsection.get("title") or "").strip()
        match = re.match(r"(\d+(?:\.\d+)*)", heading)
        prefix = f"{match.group(1)}.{index}" if match else ""
        numbered = f"{prefix} {title}" if prefix else title
        add_section(document, numbered, subsection, f"{field}.{index}", min(level + 1, 6))


def replace_text_in_runs(paragraph: Any, replacements: dict[str, str]) -> None:
    runs = list(paragraph.runs)
    if not runs:
        return
    combined = "".join(run.text or "" for run in runs)
    updated = combined
    for old, new in replacements.items():
        if old == "V1.0":
            updated = re.sub(r"(?<!\w)V1\.0(?:\.0)?(?![.\d])", new, updated)
        else:
            updated = updated.replace(old, new)
    if updated == combined:
        return
    # 封面版本号在模板里分散于多个 run；保留首个 run 的字符样式和段落格式。
    runs[0].text = updated
    for run in runs[1:]:
        run.text = ""


def replace_template_identity(document: Any, software_name: str, version: str) -> None:
    replacements = {
        "{{SOFTWARE_NAME}}": software_name,
        "V1.0": version,
    }
    for paragraph in document.paragraphs:
        replace_text_in_runs(paragraph, replacements)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_text_in_runs(paragraph, replacements)
    for section in document.sections:
        for container in (section.header, section.first_page_header, section.footer, section.first_page_footer):
            for paragraph in container.paragraphs:
                replace_text_in_runs(paragraph, replacements)


def format_design_cover(document: Any, software_name: str, version: str) -> None:
    title = next((p for p in document.paragraphs if p.text.strip() in {software_name, f"{software_name}软件说明书"}), None)
    subtitle = next((p for p in document.paragraphs if p.text.strip() == "系统设计说明书"), None)
    version_paragraph = next((p for p in document.paragraphs if p.text.strip() == version), None)
    if title is None or version_paragraph is None:
        raise RuntimeError("设计说明书模板封面缺少标题或版本号")
    title.clear()
    title_run = title.add_run(f"{software_name}软件说明书")
    title_run.font.name = "宋体"
    title_run.font.size = Pt(26)  # 一号
    title_run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    if subtitle is not None:
        subtitle._element.getparent().remove(subtitle._element)
    version_paragraph.clear()
    version_run = version_paragraph.add_run(version)
    version_run.font.name = "Times New Roman"
    version_run.font.size = Pt(22)  # 二号


def clear_template_content_after_toc(document: Any) -> None:
    body = document._element.body
    content_start = None
    for child in list(body):
        if child.tag.endswith("}p"):
            text = "".join(node.text or "" for node in child.iter() if node.tag.endswith("}t")).strip()
            if text in {"背景", "系统架构设计"}:
                content_start = child
                break
    if content_start is None:
        raise RuntimeError("设计说明书模板中未找到正文起点，不能安全保留首页和目录")
    started = False
    for child in list(body):
        if child is content_start:
            started = True
        if started and not child.tag.endswith("}sectPr"):
            body.remove(child)


def request_toc_refresh(document: Any) -> None:
    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def refresh_toc_in_word(path: Path) -> None:
    """Save real TOC page numbers when Word is available on macOS.

    Other systems retain the updateFields flag so Word can refresh the directory
    when the recipient opens the DOCX.
    """
    if platform.system() != "Darwin" or not shutil.which("osascript"):
        return
    probe = subprocess.run(
        ["osascript", "-e", 'POSIX path of (path to application id "com.microsoft.Word")'],
        capture_output=True,
        text=True,
        timeout=15,
    )
    if probe.returncode != 0:
        return
    posix_path = str(path.resolve()).replace("\\", "\\\\").replace('"', '\\"')
    doc_name = path.name.replace("\\", "\\\\").replace('"', '\\"')
    script = f'''
set docPosix to "{posix_path}"
set docName to "{doc_name}"
with timeout of 600 seconds
  tell application "Microsoft Word"
    set savedAlerts to display alerts
    set display alerts to alerts none
    set theDoc to missing value
    try
      try
        if (exists document docName) then close document docName saving no
      end try
      open (POSIX file docPosix)
      repeat 80 times
        if (exists document docName) then
          set theDoc to document docName
          exit repeat
        end if
        delay 0.5
      end repeat
      if theDoc is missing value then error "Word did not open design specification"
      repaginate theDoc
      update field (every field of text object of theDoc)
      repaginate theDoc
      save theDoc
      close theDoc saving no
      set theDoc to missing value
    on error errMsg number errNum
      try
        if theDoc is not missing value then close theDoc saving no
      end try
      set display alerts to savedAlerts
      error errMsg number errNum
    end try
    set display alerts to savedAlerts
  end tell
end timeout
'''
    result = subprocess.run(
        ["osascript", "-"],
        input=script,
        capture_output=True,
        text=True,
        timeout=660,
    )
    if result.returncode != 0:
        raise RuntimeError(f"Microsoft Word 更新设计说明书目录失败：{result.stderr.strip() or result.stdout.strip()}")


def rebuild_toc_cache(document: Any) -> None:
    """Preserve the template TOC container and formatting while removing stale entries."""
    toc = next((child for child in document._element.body if child.tag == qn("w:sdt")), None)
    if toc is None:
        raise RuntimeError("设计说明书模板缺少目录内容控件")
    content = toc.find(qn("w:sdtContent"))
    if content is None:
        raise RuntimeError("设计说明书模板目录内容不可读取")
    old_entries = [child for child in content if child.tag == qn("w:p")]
    if len(old_entries) < 2:
        raise RuntimeError("设计说明书模板目录缺少示例条目")
    entry_properties = old_entries[1].find(qn("w:pPr"))
    for child in list(content):
        if child is not old_entries[0] and child.tag != qn("w:bookmarkEnd"):
            content.remove(child)

    headings = [
        (paragraph, paragraph.style.name == "Heading 2")
        for paragraph in document.paragraphs
        if paragraph.style.name == "Heading 1"
    ]
    if not headings:
        raise RuntimeError("设计说明书正文没有可用于目录的标题")

    def add_field(paragraph: Any, field_type: str) -> None:
        run = OxmlElement("w:r")
        fld = OxmlElement("w:fldChar")
        fld.set(qn("w:fldCharType"), field_type)
        if field_type == "begin":
            fld.set(qn("w:dirty"), "true")
        run.append(fld)
        paragraph.append(run)

    for index, (heading, nested) in enumerate(headings):
        bookmark_name = f"_DesignToc{index + 1}"
        bookmark_id = str(10000 + index)
        bookmark_start = OxmlElement("w:bookmarkStart")
        bookmark_start.set(qn("w:id"), bookmark_id)
        bookmark_start.set(qn("w:name"), bookmark_name)
        heading._p.insert(1 if heading._p.find(qn("w:pPr")) is not None else 0, bookmark_start)
        bookmark_end = OxmlElement("w:bookmarkEnd")
        bookmark_end.set(qn("w:id"), bookmark_id)
        heading._p.append(bookmark_end)

        paragraph = OxmlElement("w:p")
        if entry_properties is not None:
            properties = deepcopy(entry_properties)
            if nested:
                indent = properties.find(qn("w:ind"))
                if indent is None:
                    indent = OxmlElement("w:ind")
                    properties.append(indent)
                indent.set(qn("w:left"), "360")
            paragraph.append(properties)
        if index == 0:
            add_field(paragraph, "begin")
            instruction_run = OxmlElement("w:r")
            instruction = OxmlElement("w:instrText")
            instruction.text = 'TOC \\o "1-1" \\h'
            instruction_run.append(instruction)
            paragraph.append(instruction_run)
            add_field(paragraph, "separate")
        run = OxmlElement("w:r")
        label = OxmlElement("w:t")
        label.text = heading.text
        run.append(label)
        paragraph.append(run)
        tab_run = OxmlElement("w:r")
        tab_run.append(OxmlElement("w:tab"))
        paragraph.append(tab_run)
        add_field(paragraph, "begin")
        page_instruction_run = OxmlElement("w:r")
        page_instruction = OxmlElement("w:instrText")
        page_instruction.text = f"PAGEREF {bookmark_name} \\h"
        page_instruction_run.append(page_instruction)
        paragraph.append(page_instruction_run)
        add_field(paragraph, "separate")
        page_run = OxmlElement("w:r")
        page_text = OxmlElement("w:t")
        page_text.text = "1"
        page_run.append(page_text)
        paragraph.append(page_run)
        add_field(paragraph, "end")
        if index == len(headings) - 1:
            add_field(paragraph, "end")
        content.append(paragraph)


def add_design_content(document: Any, business: dict[str, Any]) -> None:
    for chapter, (field, title) in enumerate(REQUIRED_FIELDS, start=1):
        add_section(document, f"{chapter} {title}", business[field], field)


def markdown_design(software_name: str, version: str, business: dict[str, Any]) -> str:
    lines = [f"# {software_name} V{version.lstrip('Vv')} 软件设计说明书", ""]
    for chapter, (field, title) in enumerate(REQUIRED_FIELDS, start=1):
        value = business.get(field)
        lines.extend([f"## {chapter} {title}", ""])
        lines.extend(section_markdown(value, str(chapter)))
    return "\n".join(lines) + "\n"


def write_design_draft(business: dict[str, Any], software_name: str, version: str, out_dir: Path) -> Path:
    validate_design_structure(business)
    draft_md = ensure_dir(out_dir) / "软件设计说明书.md"
    draft_md.write_text(markdown_design(software_name, version, business), encoding="utf-8")
    return draft_md


def generate_design_docx(
    business: dict[str, Any],
    software_name: str,
    version: str,
    template: Path,
    out_dir: Path,
) -> tuple[Path, Path]:
    if Document is None:
        raise RuntimeError(f"python-docx 不可用：{DOCX_IMPORT_ERROR}")
    draft_md = write_design_draft(business, software_name, version, out_dir)
    workdir = out_dir.parent
    if not template.exists():
        raise FileNotFoundError(f"缺少设计说明书模板：{template}")
    final_dir = ensure_dir(workdir / "正式资料")
    output = final_dir / f"{safe_filename(software_name)}-{version}-软件设计说明书.docx"
    shutil.copy2(template, output)
    document = Document(str(output))
    replace_template_identity(document, software_name, version)
    format_design_cover(document, software_name, version)
    clear_template_content_after_toc(document)
    document.add_page_break()
    add_design_content(document, business)
    rebuild_toc_cache(document)
    request_toc_refresh(document)
    document.core_properties.title = f"{software_name}-{version}-软件设计说明书"
    document.save(str(output))
    refresh_toc_in_word(output)
    return draft_md, output


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--business-context", required=True)
    parser.add_argument("--software-name", required=True)
    parser.add_argument("--version", default="V1.0.0")
    parser.add_argument("--template", default="")
    parser.add_argument("--out-dir", default="软件著作权申请资料/草稿")
    parser.add_argument("--draft-only", action="store_true", help="只生成供用户确认的 Markdown 草稿")
    args = parser.parse_args()
    if Document is None:
        raise SystemExit(f"STOP_FOR_USER\nNEXT_ACTION: python-docx 不可用，无法生成设计说明书：{DOCX_IMPORT_ERROR}")

    business = read_json(Path(args.business_context))
    if business.get("confirmation_required") and not business.get("user_confirmed"):
        raise SystemExit("STOP_FOR_USER\nNEXT_ACTION: 业务理解尚未确认，须先记录 business 确认点。")
    out_dir = ensure_dir(Path(args.out_dir))
    if args.draft_only:
        draft_md = write_design_draft(business, args.software_name, args.version, out_dir)
        print(f"OK design draft: {draft_md}")
        return
    template = Path(args.template) if args.template else Path(__file__).resolve().parent.parent / "assets" / "design_spec_template.docx"
    try:
        draft_md, output = generate_design_docx(business, args.software_name, args.version, template, out_dir)
    except Exception as exc:
        raise SystemExit(f"STOP_FOR_USER\nNEXT_ACTION: 软件设计说明书生成失败：{exc}") from exc
    print(f"OK design draft: {draft_md}")
    print(f"OK design docx: {output}")


if __name__ == "__main__":
    main()
